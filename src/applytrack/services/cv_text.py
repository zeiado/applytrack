"""Reads the text of an uploaded CV, for Sekka AI. Runs on the job queue.

The file came from a user, so it's checked before it's parsed: the first bytes must
match the claimed type, and a Word file may not unpack to more than a set size.
"""

import re
import zipfile
from io import BytesIO

import docx
from pypdf import PdfReader

from applytrack.schemas.cvs import DOCX, PDF

MAX_PAGES = 30
MAX_TEXT_CHARS = 50_000
MAX_UNPACKED_BYTES = 30 * 1024 * 1024
NUL = chr(0)


class UnreadableCV(Exception):  # noqa: N818 (a result, not a failure of ours)
    """The file can't be read. ``code`` is shown to the user (CVRead.error)."""

    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


def extract_text(data: bytes, content_type: str) -> str:
    if content_type == PDF and data.startswith(b"%PDF-"):
        text = _pdf_text(data)
    elif content_type == DOCX and data.startswith(b"PK\x03\x04"):
        text = _docx_text(data)
    else:
        raise UnreadableCV("wrong_type")
    text = tidy(text)
    if not text:
        raise UnreadableCV("no_text")  # e.g. a scanned page, which is only an image
    return text[:MAX_TEXT_CHARS]


def _pdf_text(data: bytes) -> str:
    try:
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(""):
            raise UnreadableCV("encrypted")
        return "\n\n".join(page.extract_text() or "" for page in reader.pages[:MAX_PAGES])
    except UnreadableCV:
        raise
    except Exception as exc:  # pypdf raises many kinds of errors for broken files
        raise UnreadableCV("unreadable") from exc


def _docx_text(data: bytes) -> str:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if "word/document.xml" not in archive.namelist():
                raise UnreadableCV("wrong_type")
            if sum(item.file_size for item in archive.infolist()) > MAX_UNPACKED_BYTES:
                raise UnreadableCV("unreadable")  # a "zip bomb"
        document = docx.Document(BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    except UnreadableCV:
        raise
    except Exception as exc:
        raise UnreadableCV("unreadable") from exc


def tidy(text: str) -> str:
    """One space between words, no blank lines in a row."""
    lines = (" ".join(line.split()) for line in text.replace(NUL, "").splitlines())
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
