"""Uploading CVs (#115): signed upload links, reading the text on the job queue, limits,
and deleting files with the CV or the account."""

import uuid
from collections.abc import Iterator
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlsplit

import docx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from applytrack.api.deps import get_now
from applytrack.core.config import Settings
from applytrack.jobs.worker import JobContext, Worker
from applytrack.models import CV, CVStatus
from applytrack.schemas.cvs import DOCX, MAX_CV_BYTES, PDF
from applytrack.services.cvs import AbandonedUploadCleaner, ExtractCVTextHandler
from applytrack.storage import StorageError
from applytrack.storage.local import LocalStorage
from tests.conftest import AuthHelper, error_without_request_id

ARABIC_NAME = "".join(chr(c) for c in (0x633, 0x64A, 0x631, 0x629)) + ".pdf"  # "سيرة.pdf"


def make_pdf(text: str | None) -> bytes:
    """A one-page PDF showing ``text`` (or a blank page, like a scanned image)."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return out


def make_docx(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    out = BytesIO()
    document.save(out)
    return out.getvalue()


@pytest.fixture
def settings(settings: Settings, tmp_path: Path) -> Settings:
    return settings.model_copy(update={"storage_local_dir": str(tmp_path / "files")})


@pytest.fixture
def storage(app: FastAPI) -> LocalStorage:
    storage: LocalStorage = app.state.storage
    return storage


@pytest.fixture
def headers(auth: AuthHelper) -> dict[str, str]:
    return auth.headers()


class CVClient:
    def __init__(self, client: TestClient, headers: dict[str, str]) -> None:
        self.client = client
        self.headers = headers

    def start(
        self, data: bytes, *, content_type: str = PDF, filename: str = "cv.pdf"
    ) -> dict[str, Any]:
        response = self.client.post(
            "/api/v1/cvs",
            json={"filename": filename, "content_type": content_type, "size_bytes": len(data)},
            headers=self.headers,
        )
        assert response.status_code == 201, response.text
        body: dict[str, Any] = response.json()
        return body

    def put(self, start: dict[str, Any], data: bytes, **headers: str) -> int:
        upload: dict[str, Any] = start["upload"]
        response = self.client.put(
            upload["url"], content=data, headers={**upload["headers"], **headers}
        )
        status: int = response.status_code
        return status

    def upload(self, data: bytes, **kwargs: str) -> dict[str, Any]:
        """Start, upload and complete: the CV is then being read."""
        start = self.start(data, **kwargs)
        assert self.put(start, data) == 200
        response = self.complete(start["cv"]["id"])
        assert response.status_code == 200, response.text
        body: dict[str, Any] = response.json()
        return body

    def complete(self, cv_id: str) -> Any:
        return self.client.post(f"/api/v1/cvs/{cv_id}/complete", headers=self.headers)

    def get(self, cv_id: str) -> dict[str, Any]:
        response = self.client.get(f"/api/v1/cvs/{cv_id}", headers=self.headers)
        assert response.status_code == 200, response.text
        body: dict[str, Any] = response.json()
        return body

    def list(self) -> list[dict[str, Any]]:
        body: list[dict[str, Any]] = self.client.get("/api/v1/cvs", headers=self.headers).json()
        return body


@pytest.fixture
def cvs(client: TestClient, headers: dict[str, str]) -> CVClient:
    return CVClient(client, headers)


@pytest.fixture
def frozen_now(app: FastAPI) -> Iterator[list[datetime]]:
    """Set ``frozen_now[0]`` to move the API's clock."""
    now = [datetime(2026, 10, 26, 9, tzinfo=UTC)]
    app.dependency_overrides[get_now] = lambda: now[0]
    yield now
    app.dependency_overrides.pop(get_now)


def stored_files(storage: LocalStorage) -> list[str]:
    return sorted(p.name for p in storage.directory.rglob("*") if p.is_file())


# --- the happy paths ---


def test_a_pdf_is_uploaded_and_its_text_read(
    cvs: CVClient, worker: Worker, storage: LocalStorage
) -> None:
    pdf = make_pdf("Python and SQL developer")
    start = cvs.start(pdf, filename="Ada Lovelace CV.pdf")
    assert start["cv"]["status"] == "uploading"
    assert start["upload"]["method"] == "PUT"
    assert start["upload"]["headers"] == {"Content-Type": PDF}
    assert start["upload"]["url"].startswith("http://localhost:8000/local-storage/cvs/")

    assert cvs.put(start, pdf) == 200
    completed = cvs.complete(start["cv"]["id"]).json()
    assert completed["status"] == "processing"

    worker.run_pending()
    cv = cvs.get(start["cv"]["id"])
    assert cv["status"] == "ready"
    assert cv["error"] is None
    assert "Python and SQL developer" in cv["text"]
    assert cv["filename"] == "Ada Lovelace CV.pdf"
    assert cv["size_bytes"] == len(pdf)
    assert [c["id"] for c in cvs.list()] == [cv["id"]]
    assert "text" not in cvs.list()[0]  # the list stays small
    assert len(stored_files(storage)) == 1


def test_a_word_file_is_read_with_its_tables(cvs: CVClient, worker: Worker) -> None:
    data = make_docx("Ada Lovelace", "  Analyst  ", table=[["Skills", "Excel"]])
    cv = cvs.upload(data, content_type=DOCX, filename="cv.docx")
    worker.run_pending()
    text = cvs.get(cv["id"])["text"]
    assert text == "Ada Lovelace\nAnalyst\nSkills | Excel"


def test_a_cv_downloads_under_its_own_name(
    client: TestClient, cvs: CVClient, headers: dict[str, str]
) -> None:
    pdf = make_pdf("hello")
    cv = cvs.upload(pdf, filename=ARABIC_NAME)
    response = client.get(f"/api/v1/cvs/{cv['id']}/download", headers=headers)
    assert response.status_code == 200
    file = client.get(response.json()["url"])
    assert file.status_code == 200
    assert file.content == pdf
    assert file.headers["content-disposition"] == (
        "attachment; filename=\"____.pdf\"; filename*=UTF-8''%D8%B3%D9%8A%D8%B1%D8%A9.pdf"
    )


def test_the_filename_is_kept_without_folders(cvs: CVClient) -> None:
    start = cvs.start(b"x", filename="C:\\Users\\ada\\cv final.pdf")
    assert start["cv"]["filename"] == "cv final.pdf"


# --- files that can't be read ---


@pytest.mark.parametrize(
    ("data", "content_type", "error"),
    [
        (b"just some text", PDF, "wrong_type"),
        (make_pdf(None), PDF, "no_text"),
        (b"%PDF-1.4 broken", PDF, "unreadable"),
        (make_pdf("hi"), DOCX, "wrong_type"),
    ],
)
def test_files_that_cant_be_read_fail_with_a_reason(
    cvs: CVClient, worker: Worker, data: bytes, content_type: str, error: str
) -> None:
    cv = cvs.upload(data, content_type=content_type)
    worker.run_pending()
    read = cvs.get(cv["id"])
    assert (read["status"], read["error"], read["text"]) == ("failed", error, None)


# --- what the upload link accepts ---


def test_the_upload_link_only_takes_the_announced_file(
    cvs: CVClient, storage: LocalStorage
) -> None:
    pdf = make_pdf("hello")
    start = cvs.start(pdf)
    assert cvs.put(start, pdf + b"more") == 400
    assert cvs.put(start, pdf, **{"Content-Type": "text/html"}) == 403
    url = start["upload"]["url"].replace("signature=", "signature=0")
    assert cvs.put({"upload": {**start["upload"], "url": url}}, pdf) == 403
    assert stored_files(storage) == []


def test_an_expired_upload_link_is_refused(cvs: CVClient, storage: LocalStorage) -> None:
    pdf = make_pdf("hello")
    start = cvs.start(pdf)
    storage.clock = lambda: datetime.now(UTC).timestamp() + 3600
    assert cvs.put(start, pdf) == 403


def test_uploads_answer_the_web_apps_cors_preflight(
    client: TestClient, cvs: CVClient, settings: Settings
) -> None:
    start = cvs.start(b"x")
    response = client.options(start["upload"]["url"])
    assert response.status_code == 204
    assert response.headers["access-control-allow-origin"] == settings.public_app_url
    assert response.headers["access-control-allow-methods"] == "PUT"


def test_completing_before_the_upload_finished_is_refused(cvs: CVClient) -> None:
    start = cvs.start(b"data")
    response = cvs.complete(start["cv"]["id"])
    assert response.status_code == 422
    assert error_without_request_id(response)["error"]["code"] == "upload_missing"


def test_completing_twice_is_refused(cvs: CVClient) -> None:
    cv = cvs.upload(make_pdf("hello"))
    response = cvs.complete(cv["id"])
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "already_uploaded"


def test_a_file_of_another_size_is_dropped(cvs: CVClient, storage: LocalStorage) -> None:
    start = cvs.start(b"12345")
    storage.write(start["upload"]["url"].split("/local-storage/")[1].split("?")[0], b"123")
    response = cvs.complete(start["cv"]["id"])
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "upload_invalid"
    assert cvs.list() == []


@pytest.mark.parametrize(
    "body",
    [
        {"filename": "cv.pdf", "content_type": PDF, "size_bytes": MAX_CV_BYTES + 1},
        {"filename": "cv.pdf", "content_type": PDF, "size_bytes": 0},
        {"filename": "cv.png", "content_type": "image/png", "size_bytes": 10},
        {"filename": " / ", "content_type": PDF, "size_bytes": 10},
        {"filename": "cv.pdf", "content_type": PDF, "size_bytes": 10, "extra": 1},
    ],
)
def test_upload_requests_are_validated(
    client: TestClient, headers: dict[str, str], body: dict[str, Any]
) -> None:
    response = client.post("/api/v1/cvs", json=body, headers=headers)
    assert response.status_code == 422


# --- limits ---


def test_up_to_three_cvs_and_expired_uploads_dont_count(
    cvs: CVClient, frozen_now: list[datetime]
) -> None:
    for _ in range(3):
        cvs.start(b"x")
    response = cvs.client.post(
        "/api/v1/cvs",
        json={"filename": "cv.pdf", "content_type": PDF, "size_bytes": 1},
        headers=cvs.headers,
    )
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "too_many_cvs"

    frozen_now[0] += timedelta(minutes=11)  # the three links have expired
    cvs.start(b"x")
    assert len(cvs.list()) == 1


# --- ownership, deletion ---


def test_other_peoples_cvs_are_not_found(
    client: TestClient, cvs: CVClient, auth: AuthHelper
) -> None:
    cv = cvs.upload(make_pdf("hello"))
    other = auth.headers("grace@example.com")
    for method, path in [
        ("GET", ""),
        ("GET", "/download"),
        ("POST", "/complete"),
        ("DELETE", ""),
    ]:
        response = client.request(method, f"/api/v1/cvs/{cv['id']}{path}", headers=other)
        assert response.status_code == 404, (method, path)
        assert response.json()["error"]["code"] == "cv_not_found"


def test_deleting_a_cv_deletes_its_file(
    client: TestClient, cvs: CVClient, worker: Worker, storage: LocalStorage
) -> None:
    cv = cvs.upload(make_pdf("hello"))
    response = client.delete(f"/api/v1/cvs/{cv['id']}", headers=cvs.headers)
    assert response.status_code == 204
    assert cvs.list() == []
    worker.run_pending()
    assert stored_files(storage) == []


def test_deleting_the_account_deletes_the_files(
    client: TestClient, cvs: CVClient, worker: Worker, storage: LocalStorage
) -> None:
    cvs.upload(make_pdf("one"))
    cvs.upload(make_pdf("two"))
    worker.run_pending()
    assert len(stored_files(storage)) == 2
    response = client.request(
        "DELETE", "/api/v1/users/me", json={"password": "correct-horse"}, headers=cvs.headers
    )
    assert response.status_code == 204
    worker.run_pending()
    assert stored_files(storage) == []


def test_the_export_includes_cvs_and_their_text(
    client: TestClient, cvs: CVClient, worker: Worker
) -> None:
    cvs.upload(make_pdf("Data analyst"), filename="mine.pdf")
    worker.run_pending()
    export = client.get("/api/v1/users/me/export", headers=cvs.headers).json()
    [cv] = export["cvs"]
    assert cv["filename"] == "mine.pdf"
    assert "Data analyst" in cv["text"]


def test_downloading_before_the_upload_is_refused(cvs: CVClient) -> None:
    start = cvs.start(b"x")
    response = cvs.client.get(f"/api/v1/cvs/{start['cv']['id']}/download", headers=cvs.headers)
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "not_uploaded"


def test_abandoned_uploads_are_cleaned_up(
    app: FastAPI, cvs: CVClient, worker: Worker, storage: LocalStorage
) -> None:
    start = cvs.start(b"x")
    storage.write(start["upload"]["url"].split("/local-storage/")[1].split("?")[0], b"x")
    done = cvs.upload(make_pdf("kept"))
    cleaner = AbandonedUploadCleaner(app.state.session_factory)
    now = datetime.now(UTC)
    assert cleaner.run(now) == 0  # too recent
    assert cleaner.run(now + timedelta(hours=2)) == 1
    worker.run_pending()
    assert [c["id"] for c in cvs.list()] == [done["id"]]
    assert len(stored_files(storage)) == 1


# --- storage trouble ---


def test_uploads_are_unavailable_without_storage(
    app: FastAPI, client: TestClient, headers: dict[str, str]
) -> None:
    app.state.storage = None
    response = client.post(
        "/api/v1/cvs",
        json={"filename": "cv.pdf", "content_type": PDF, "size_bytes": 1},
        headers=headers,
    )
    assert response.status_code == 503
    assert response.json()["error"]["code"] == "uploads_unavailable"
    assert client.get("/api/v1/cvs", headers=headers).json() == []


def test_storage_down_while_completing(
    cvs: CVClient, storage: LocalStorage, monkeypatch: pytest.MonkeyPatch
) -> None:
    start = cvs.start(b"x")

    def down(_: str) -> int:
        raise StorageError("timeout")

    monkeypatch.setattr(storage, "size", down)
    response = cvs.complete(start["cv"]["id"])
    assert response.status_code == 503
    assert response.json()["error"]["code"] == "storage_unavailable"


class _DownStorage(LocalStorage):
    def read(self, key: str) -> bytes:
        raise StorageError("timeout")


def test_reading_retries_then_gives_up_on_storage_errors(
    app: FastAPI, cvs: CVClient, storage: LocalStorage
) -> None:
    cv = cvs.upload(make_pdf("hello"))
    down = _DownStorage(storage.directory, public_url="http://x", secret="s")
    handler = ExtractCVTextHandler(app.state.session_factory, down)
    payload = {"cv_id": cv["id"]}
    with pytest.raises(StorageError):
        handler(payload, JobContext(uuid.uuid4(), "extract_cv_text", 1, 3))
    assert cvs.get(cv["id"])["status"] == "processing"  # will retry
    with pytest.raises(StorageError):
        handler(payload, JobContext(uuid.uuid4(), "extract_cv_text", 3, 3))
    read = cvs.get(cv["id"])
    assert (read["status"], read["error"]) == ("failed", "storage_error")


def test_reading_a_deleted_cv_does_nothing(
    app: FastAPI, cvs: CVClient, storage: LocalStorage
) -> None:
    handler = ExtractCVTextHandler(app.state.session_factory, storage)
    handler({"cv_id": str(uuid.uuid4())}, JobContext(uuid.uuid4(), "extract_cv_text", 1, 3))
    with app.state.session_factory() as session:
        assert session.query(CV).count() == 0
    assert CVStatus.READY.value == "ready"


def test_the_local_link_carries_its_signature(storage: LocalStorage) -> None:
    link = storage.upload_link("cvs/abc", content_type=PDF, size=10, expires=timedelta(minutes=1))
    query = parse_qs(urlsplit(link.url).query)
    assert set(query) == {"type", "size", "expires", "signature"}
    with pytest.raises(ValueError, match="Invalid storage key"):
        storage.path("../etc/passwd")


# --- replacing ---


def test_replacing_a_cv_works_at_the_limit(
    cvs: CVClient, worker: Worker, storage: LocalStorage
) -> None:
    first = cvs.upload(make_pdf("old"))
    cvs.upload(make_pdf("two"))
    cvs.upload(make_pdf("three"))
    worker.run_pending()
    new_pdf = make_pdf("new")
    response = cvs.client.post(
        "/api/v1/cvs",
        json={
            "filename": "new.pdf",
            "content_type": PDF,
            "size_bytes": len(new_pdf),
            "replaces": first["id"],
        },
        headers=cvs.headers,
    )
    assert response.status_code == 201, response.text
    start = response.json()
    assert len(cvs.list()) == 4  # the old one stays until the new one is uploaded

    # Only one replacement at a time: the limit still holds.
    again = cvs.client.post(
        "/api/v1/cvs",
        json={"filename": "x.pdf", "content_type": PDF, "size_bytes": 1, "replaces": first["id"]},
        headers=cvs.headers,
    )
    assert again.status_code == 409

    assert cvs.put(start, new_pdf) == 200
    assert cvs.complete(start["cv"]["id"]).status_code == 200
    worker.run_pending()
    names = [c["filename"] for c in cvs.list()]
    assert names == ["cv.pdf", "cv.pdf", "new.pdf"]
    assert first["id"] not in [c["id"] for c in cvs.list()]
    assert len(stored_files(storage)) == 3


def test_only_your_own_cv_can_be_replaced(cvs: CVClient, auth: AuthHelper) -> None:
    cv = cvs.upload(make_pdf("mine"))
    other = auth.headers("grace@example.com")
    response = cvs.client.post(
        "/api/v1/cvs",
        json={"filename": "x.pdf", "content_type": PDF, "size_bytes": 1, "replaces": cv["id"]},
        headers=other,
    )
    assert response.status_code == 404
